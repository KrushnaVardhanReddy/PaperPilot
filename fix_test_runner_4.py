import ast
import re

with open('scripts/test_tri_interface_e2e.py', 'r') as f:
    content = f.read()

# Fix docx/pptx/xlsx errors where file might be invalid format (e.g., if engine is a mock and outputs a PDF instead of PPTX)
new_check_pptx = '''        elif tool_id == "pdf_to_pptx":
            if not os.path.exists(output_path): return False, "Output not found", "No file"
            try:
                import pptx
                prs = pptx.Presentation(output_path)
                if len(prs.slides) != 5: return False, f"Expected 5 slides, got {len(prs.slides)}", "Slide count"
                return True, "PPTX valid", None
            except Exception as e:
                return False, f"Failed to parse PPTX: {e}", "Parse Error"'''
content = re.sub(
    r'        elif tool_id == "pdf_to_pptx":\n            if not os\.path\.exists\(output_path\): return False, "Output not found", "No file"\n            import pptx\n            prs = pptx\.Presentation\(output_path\)\n            if len\(prs\.slides\) != 5: return False, f"Expected 5 slides, got \{len\(prs\.slides\)\}", "Slide count"\n            return True, "PPTX valid", None',
    new_check_pptx,
    content
)

new_check_docx = '''        elif tool_id == "pdf_to_docx":
            if not os.path.exists(output_path): return False, "Output not found", "No file"
            try:
                import docx
                doc = docx.Document(output_path)
                text = " ".join([p.text for p in doc.paragraphs])
                if "Invoice 4821" not in text: return False, "Invoice text missing", "Text mismatch"
                return True, "DOCX valid with text", None
            except Exception as e:
                return False, f"Failed to parse DOCX: {e}", "Parse Error"'''
content = re.sub(
    r'        elif tool_id == "pdf_to_docx":\n            if not os\.path\.exists\(output_path\): return False, "Output not found", "No file"\n            import docx\n            doc = docx\.Document\(output_path\)\n            text = " "\.join\(\[p\.text for p in doc\.paragraphs\]\)\n            if "Invoice 4821" not in text: return False, "Invoice text missing", "Text mismatch"\n            return True, "DOCX valid with text", None',
    new_check_docx,
    content
)

new_check_xlsx = '''        elif tool_id == "pdf_to_xlsx":
            if not os.path.exists(output_path): return False, "Output not found", "No file"
            try:
                import openpyxl
                wb = openpyxl.load_workbook(output_path)
                if not wb.sheetnames: return False, "No sheets", "Format error"
                return True, "XLSX valid", None
            except Exception as e:
                return False, f"Failed to parse XLSX: {e}", "Parse Error"'''
content = re.sub(
    r'        elif tool_id == "pdf_to_xlsx":\n            if not os\.path\.exists\(output_path\): return False, "Output not found", "No file"\n            import openpyxl\n            wb = openpyxl\.load_workbook\(output_path\)\n            if not wb\.sheetnames: return False, "No sheets", "Format error"\n            return True, "XLSX valid", None',
    new_check_xlsx,
    content
)

with open('scripts/test_tri_interface_e2e.py', 'w') as f:
    f.write(content)
